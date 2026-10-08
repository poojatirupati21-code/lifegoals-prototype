"""Helper for calc-vs-xlsx.js: write each round's inputs into a copy of the workbook (every sheet at once),
recalculate with LibreOffice (recalc.py), and read the named results with openpyxl (data_only).
Input JSON (argv[1]): {"rounds": [{"Sheet name": {"Name": value | null, ...}, ...}, ...], "read": {"Sheet name": ["Name", ...]}, "assumptions": ["Name", ...]}
Output JSON (argv[2]): {"rounds": [{"Sheet name": {"Name": value}}], "assumptions": {"Name": value}, "mtime": ...}"""
import json, os, shutil, subprocess, sys, datetime
import openpyxl
from openpyxl.utils.cell import range_boundaries

HERE = os.path.dirname(os.path.abspath(__file__))
WB = os.environ.get('WORKBOOK', os.path.join(HERE, '..', '..', 'deliverables', 'LifeGoals-Calculators.xlsx'))
RECALC = os.environ.get('RECALC', '/root/.claude/skills/synced/849805a0-1896-4a49-87bc-2f01ccbb7dae_a4e27067-8324-46de-8d01-0f71917b608e/xlsx/scripts/recalc.py')
OUT = os.path.join(os.environ.get('CALC_VS_XLSX_WORK', os.path.join(__import__('tempfile').gettempdir(), 'calc-vs-xlsx')), 'xlsxrounds')


def dest(wb, ws, name):
    d = ws.defined_names.get(name) if hasattr(ws, 'defined_names') else None
    if d is None:
        d = wb.defined_names.get(name)
    if d is None:
        return None
    for sheet, ref in d.destinations:
        c1, r1, c2, r2 = range_boundaries(ref.replace('$', ''))
        return wb[sheet], c1, r1, c2, r2
    return None


def main():
    spec = json.load(open(sys.argv[1]))
    os.makedirs(OUT, exist_ok=True)
    mtime = os.path.getmtime(WB)
    res = {'rounds': [], 'mtime': mtime}
    for k, rnd in enumerate(spec['rounds']):
        path = '%s/round%d.xlsx' % (OUT, k)
        shutil.copy(WB, path)
        wb = openpyxl.load_workbook(path)
        for sheet, vals in rnd.items():
            ws = wb[sheet]
            for name, v in vals.items():
                d = dest(wb, ws, name + '_Entry') or dest(wb, ws, name)   # frozen §14 layout: the customer's figure goes in the *_Entry cell
                if d is None:
                    raise SystemExit('missing name %s on %s' % (name, sheet))
                tws, c1, r1, _, _ = d
                cell = tws.cell(row=r1, column=c1)
                if isinstance(v, str) and v.startswith('DATE:'):
                    v = datetime.datetime.strptime(v[5:], '%Y-%m-%d')
                cell.value = v
        wb.save(path)
        r = subprocess.run([sys.executable, RECALC, path, '120'], cwd=os.path.dirname(RECALC), capture_output=True, text=True)
        if r.returncode != 0:
            raise SystemExit('recalc failed: ' + r.stdout + r.stderr)
        wv = openpyxl.load_workbook(path, data_only=True)
        out = {}
        for sheet, names in spec['read'].items():
            ws = wv[sheet]
            o = {}
            for name in names:
                d = dest(wv, ws, name)
                if d is None:
                    o[name] = '#NONAME'
                    continue
                tws, c1, r1, _, _ = d
                v = tws.cell(row=r1, column=c1).value
                o[name] = v.isoformat() if hasattr(v, 'isoformat') else v
            out[sheet] = o
        res['rounds'].append(out)
    wa = openpyxl.load_workbook(spec.get('base', WB), data_only=True)
    res['assumptions'] = {}
    for name in spec.get('assumptions', []):
        d = wa.defined_names.get(name)
        if d is None:
            res['assumptions'][name] = '#NONAME'
            continue
        vals = []
        for sheet, ref in d.destinations:
            c1, r1, c2, r2 = range_boundaries(ref.replace('$', ''))
            for row in range(r1, r2 + 1):
                vals.append(wa[sheet].cell(row=row, column=c1).value)
        res['assumptions'][name] = vals[0] if len(vals) == 1 else vals
    json.dump(res, open(sys.argv[2], 'w'), default=str)


main()
