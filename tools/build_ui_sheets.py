#!/usr/bin/env python3
"""Append the UI/UX layer to a copy of LifeGoals-Calculators.xlsx.

    python3 tools/build_ui_sheets.py <workbook.xlsx> [--out other.xlsx] [--extract ui-extract.json] [--shots dir]
                                     [--no-recalc] [--recalc path/to/recalc.py] [--keep-work dir]

Appends four sheets named "UI ..." (any existing "UI ..." sheets are replaced, so it can be re-run on the final workbook):
  UI Guide          how to read the layer, colour legend, screen order, links to every sheet
  UI Calculators    one row per screen element of all 28 calculators, from the LIVE prototype (Playwright), with the developer cell each maps to
  UI Screens        a 390 px screenshot of each calculator without a plan and with the sample plan, next to its row range
  UI Make my plan   one row per element of the plan-builder steps, the results, and the Settings standards as the customer sees them
Nothing is typed by hand: tools/extract_ui.js reads the running prototype. The developer sheets are not edited.
The workbook is saved with openpyxl (cached values are dropped) and then recalculated with recalc.py (LibreOffice) so cached values return.
"""
import argparse, json, os, re, shutil, subprocess, sys, tempfile, math
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XlImage
from openpyxl.worksheet.hyperlink import Hyperlink
from PIL import Image as PILImage

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
DEFAULT_RECALC = '/root/.claude/skills/synced/849805a0-1896-4a49-87bc-2f01ccbb7dae_a4e27067-8324-46de-8d01-0f71917b608e/xlsx/scripts/recalc.py'

TEAL, INK = '0B7A70', '0B2545'
FILL = {
    'input': 'FFF2CC', 'result': 'E2EFDA', 'gate': 'FCE4D6', 'note': 'DDEBF7', 'button': 'EDEDED', 'head': 'D9E1F2', 'none': None}
LEGEND = [
    ('Input slider', 'input', 'A figure the customer sets with a slider and a typed box (min, max and step are shown).'),
    ('Input text', 'input', 'A box the customer types in.'),
    ('Input choice', 'input', 'A Yes/No pair, chips, ticks or a list: the options are in the "Choice / chip" column.'),
    ('Gate message', 'gate', 'What the result card says until a choice is made ("Choose your ... to see this").'),
    ('Result / Result line / Sub-result row', 'result', 'The big result, the sentence under it and the small rows below it.'),
    ('Note / Tip / Example-figures label / Disclaimer', 'note', 'Text that explains, with no input.'),
    ('Button / Link', 'button', 'Something the customer taps; the behaviour is in the "Help / behaviour" column.'),
    ('Screen state', 'head', 'First row of a journey-spec 27 state in UI Make my plan: what the state is and what to look at. The rows under it are what the screen shows in that state.'),
]
thin = Side(style='thin', color='C9D3DC')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(hexc):
    return PatternFill('solid', start_color=hexc, end_color=hexc) if hexc else PatternFill(fill_type=None)


def kind_fill(t):
    t = (t or '')
    if t.startswith('Input'): return fill(FILL['input'])
    if t.startswith('Result') or t.startswith('Sub-result'): return fill(FILL['result'])
    if t.startswith('Gate'): return fill(FILL['gate'])
    if t in ('Note', 'Tip', 'Disclaimer', 'Example-figures label', 'Example-figures label (on result)', 'Banner', 'Help text', 'Eyebrow'): return fill(FILL['note'])
    if t.startswith('Button') or t == 'Link': return fill(FILL['button'])
    if t in ('Heading', 'Section header (opens)', 'Screen state'): return fill(FILL['head'])
    return fill(None)


def num(s):
    if s is None: return None
    if isinstance(s, (int, float)): return float(s)
    m = re.search(r'-?\d[\d,]*\.?\d*|-?\.\d+', str(s))
    if not m: return None
    try: return float(m.group(0).replace(',', ''))
    except ValueError: return None


def norm(s):
    return re.sub(r'[^a-z0-9]+', ' ', str(s or '').lower()).strip()


def clean(s):
    """Excel rejects control characters."""
    if s is None: return None
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', s) if isinstance(s, str) else s


# ---------------------------------------------------------------- developer workbook lookups
class Dev:
    def __init__(self, wb, wbv, calcmap):
        self.wb, self.wbv, self.map = wb, wbv, {m['id']: m for m in calcmap}
        self.names = {}  # sheet title -> name -> coordinate
        for ws in wb.worksheets:
            d = {}
            try:
                items = ws.defined_names.items()
            except Exception:
                items = []
            for n, dn in items:
                for sh, ref in dn.destinations:
                    if sh == ws.title: d[n] = ref.replace('$', '')
            self.names[ws.title] = d

    def sheet_for(self, cid, n):
        m = self.map.get(cid)
        if m and m['sheet'] in self.wb.sheetnames: return m['sheet']
        for t in self.wb.sheetnames:
            if t.startswith(n + ' '): return t
        return None

    def info(self, sheet, name):
        """Developer cell facts for a named input: effective cell, entry cell, label, unit, default (cached), allowed text, validation range."""
        co = self.names.get(sheet, {}).get(name)
        if not co: return None
        ws, wsv = self.wb[sheet], self.wbv[sheet]
        row = int(re.sub(r'\D', '', co))
        entry_co = self.names[sheet].get(name + '_Entry') or ('C%d' % row)
        dv = None
        for v in ws.data_validations.dataValidation if ws.data_validations else []:
            if entry_co in v.sqref:
                dv = v; break
        out = {'name': name, 'sheet': sheet, 'cell': co, 'entry': entry_co, 'label': ws['B%d' % row].value, 'unit': ws['D%d' % row].value,
               'allowedText': wsv['E%d' % row].value, 'default': wsv['F%d' % row].value, 'defaultFormula': ws['F%d' % row].value,
               'dvType': dv.type if dv else None, 'dvLo': None, 'dvHi': None}
        if dv and dv.type in ('decimal', 'whole') and dv.formula1 is not None:
            try:
                out['dvLo'], out['dvHi'] = float(dv.formula1), float(dv.formula2)
            except (TypeError, ValueError):
                pass
        if dv and dv.type == 'list': out['dvList'] = dv.formula1
        return out

    def result_labels(self, sheet, names):
        out = {}
        for nm in names:
            co = self.names.get(sheet, {}).get(nm)
            if co:
                row = int(re.sub(r'\D', '', co)); out[nm] = (self.wb[sheet]['B%d' % row].value, co)
        return out


def merge_fields(rec):
    """Fields of one calculator, merged across the three states, in screen order."""
    def key(f):
        sec = re.sub(r'\s*·\s*\d+ to choose$', '', f.get('section') or '')
        return (sec, f.get('label'))
    states = [('plan', rec['plan']['fields']), ('chosen', rec['chosen']['fields']), ('none', rec['none']['fields'])]
    order, byk = [], {}
    for sname, fields in states:
        prev = None
        for f in fields:
            k = key(f)
            if k not in byk:
                byk[k] = {}
                if prev is not None and prev in order: order.insert(order.index(prev) + 1, k)
                else: order.append(k)
            byk[k][sname] = f
            prev = k
    return [(k, byk[k]) for k in order]


def show(f, kind_hint=None):
    """The value as the screen shows it."""
    if f is None: return None
    if f['kind'] == 'Input choice':
        return ', '.join(f.get('selected') or []) or '(nothing selected)'
    v = f.get('value')
    if f.get('valuetext'): return f['valuetext']
    if v in (None, ''): return '(blank: the customer chooses)'
    u = f.get('unit') or ''
    return (u + v) if f.get('unitPos') == 'before' else (v + (' ' + u if u else ''))


def uniq(seq):
    out = []
    for x in seq:
        if x and x not in out: out.append(x)
    return out


def dev_compare(f, dev, kind):
    """Short verdict: does the screen's start value and slider range agree with the developer cell?"""
    if not dev: return ''
    msgs = []
    sc = 100.0 if kind == 'pct' else 1.0
    d = dev['default']
    ui = None
    if f is not None:
        if f['kind'] == 'Input choice':
            sel = (f.get('selected') or [''])[0]
            ui = 1.0 if sel == 'Yes' else 0.0 if sel == 'No' else num(sel)
        else:
            ui = num(f.get('value')) if f.get('value') not in (None, '') else None
    if ui is None: msgs.append('screen starts blank (customer chooses); developer default %s' % fmtnum(d, sc))
    elif d is None or d == '': msgs.append('developer default blank; screen %s' % fmtnum(ui, 1))
    else:
        try:
            if abs(float(d) * sc - ui) < 1e-6 * max(1, abs(ui)): msgs.append('default same (%s)' % fmtnum(ui, 1))
            else: msgs.append('DEFAULT DIFFERS: screen %s, developer %s' % (fmtnum(ui, 1), fmtnum(d, sc)))
        except (TypeError, ValueError):
            msgs.append('developer default is text: %s' % d)
    if f is not None and f.get('min') not in (None, '') and dev.get('dvLo') is not None:
        lo, hi = dev['dvLo'] * sc, dev['dvHi'] * sc
        mn, mx = num(f['min']), num(f['max'])
        if mn is not None and mx is not None:
            msgs.append('slider %s to %s is inside the developer range' % (fmtnum(mn, 1), fmtnum(mx, 1)) if lo - 1e-9 <= mn and mx <= hi + 1e-9
                        else 'SLIDER %s to %s IS OUTSIDE developer range %s to %s' % (fmtnum(mn, 1), fmtnum(mx, 1), fmtnum(lo, 1), fmtnum(hi, 1)))
    return '; '.join(msgs)


def tonum(v):
    if v in (None, ''): return v
    try:
        x = float(v)
        return int(x) if x == int(x) else x
    except (TypeError, ValueError):
        return v


def dev_shown(di, kind):
    """Developer default and allowed range as a person reads them (percent cells shown in percent)."""
    if not di: return '', ''
    sc = 100.0 if kind == 'pct' else 1.0
    d = di['default']
    try:
        dd = ('%s%%' % fmtnum(d, sc)) if kind == 'pct' and d not in (None, '') else d
    except Exception:
        dd = d
    rng = ''
    if di.get('dvLo') is not None:
        rng = '%s to %s%s' % (fmtnum(di['dvLo'], sc), fmtnum(di['dvHi'], sc), ' %' if kind == 'pct' else '')
    t = di.get('allowedText') or ''
    m = re.search(r'steps? of ([\d.,]+)', t)
    if m and rng: rng += ' · steps of ' + m.group(1)
    return dd, (rng or t)


def fmtnum(v, sc):
    try:
        x = float(v) * sc
    except (TypeError, ValueError):
        return str(v)
    return ('%.6g' % x)


# ---------------------------------------------------------------- sheet writers
HEAD_FONT = Font(bold=True, color='FFFFFF')


def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = HEAD_FONT; c.fill = fill(TEAL); c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BORDER
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, row, vals, kind=None, wrap=True):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=clean(v))
        c.alignment = Alignment(wrap_text=wrap, vertical='top'); c.border = BORDER
    if kind is not None:
        ws.cell(row=row, column=kind[0]).fill = kind[1]


def link(cell, sheet, ref='A1', text=None):
    cell.value = text or sheet
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location="'%s'!%s" % (sheet, ref), display=text or sheet)
    cell.font = Font(color='0563C1', underline='single')


def calc_rows(rec, dev, num_names):
    """Rows for one calculator: list of dicts."""
    n = rec['n']
    rows = []
    P, N, C = rec['plan'], rec['none'], rec['chosen']
    cm = dev.map.get(rec['id'], {})
    sheet = dev.sheet_for(rec['id'], n)
    inp_map = {i['key']: i for i in cm.get('inputs', [])}
    outs = cm.get('outputs', [])
    rl = dev.result_labels(sheet, outs) if sheet else {}
    def add(**kw):
        d = dict(typ='', sec='', label='', unit='', mn='', mx='', st='', noplan='', plan='', source='', chip='', help='', tags='', dname='', dcell='', ddef='', dallowed='', check='', behaviour='')
        d.update(kw); rows.append(d)
    # screen furniture
    exl = P.get('exlab') or N.get('exlab')
    if exl:
        add(typ='Example-figures label', sec='Top of the screen', label=(N.get('exlab') or {}).get('text') or (P.get('exlab') or {}).get('text') or exl['text'], noplan='Shown (no plan)' if (N.get('exlab') or {}).get('shown') else 'Hidden', plan='Hidden' if not (P.get('exlab') or {}).get('shown') else 'Shown',
            behaviour='Shown while any input still holds an example figure; hidden once the customer has their own figures or a plan (journey-spec §26).')
    if N.get('tip') or rec.get('tipDef'): add(typ='Tip', sec='Under the title', label=N.get('tip') or rec.get('tipDef'), behaviour='Light-bulb tip in the header.')
    if N.get('eyebrow'): add(typ='Note', sec='Top of the screen', label=N['eyebrow'], behaviour='Category line (the Explore group).')
    adj = N.get('adj') or P.get('adj')
    if adj: add(typ='Input choice (switch)', sec='Top of the screen', label=adj['label'], noplan='Off', plan='Off', help=adj.get('help') or '', behaviour='Switch: show the result in today\'s money. Off by default.')
    # inputs
    sigs = {}
    for (sec, label), st in merge_fields(rec):
        fp, fc, fn = st.get('plan'), st.get('chosen'), st.get('none')
        base = fp or fc or fn
        keyk = base.get('key')
        dfn = next((i for i in rec['def']['inputs'] + rec['def']['wi'] if i['k'] == keyk), None)
        mn, mx, stp = base.get('min'), base.get('max'), base.get('step')
        src = 'on screen'
        if mn in (None, '') and dfn and base['kind'] != 'Input choice':
            mn, mx, stp, src = dfn['min'], dfn['max'], dfn['step'], 'tool design (the slider appears once there is a value)'
        mp = inp_map.get(keyk)
        di = dev.info(sheet, mp['name']) if (mp and sheet) else None
        a = rec['src']['assume'].get(keyk)
        via = rec['src']['via'].get(keyk)
        pre = rec['src']['pre'].get(keyk)
        if via: source = 'Your finances / plan: ' + '; '.join(via)
        elif a: source = 'Your assumption (a choice made in Your assumptions): ' + a
        elif fp is not None and fp.get('value') not in (None, '') and dfn is not None and str(num(fp['value'])) == str(num(dfn['v'])): source = 'The tool\'s own example figure (the plan has nothing for it)'
        elif fp is not None: source = 'Customer\'s choice or the tool\'s own figure'
        else: source = ''
        chips = uniq([c['t'] for s in (fp, fc, fn) if s for c in s.get('chips', [])] + [c['t'] for s in (fp, fc, fn) if s for c in (s.get('options') if base['kind'] == 'Input choice' else [])])
        helps = uniq([h for s in (fp, fc, fn) if s for h in s.get('help', []) + s.get('hint', [])])
        tags = uniq([t for s in (fn, fc, fp) if s for t in s.get('tags', [])])
        chk = dev_compare(fn if fn is not None else fc, di, mp['kind'] if mp else 'n') if di else ('no matching cell in the developer sheet' if sheet and not mp else '')
        dd, da = dev_shown(di, mp['kind'] if mp else 'n')
        add(typ=base['kind'], sec=re.sub(r'\s*·\s*\d+ to choose$', '', sec), label=label, unit=base.get('unit') or '', mn=tonum(mn), mx=tonum(mx), st=tonum(stp), noplan=show(fn if fn is not None else fc), plan=show(fp) if fp is not None else '(not shown with a plan)',
            source=source + ('' if src == 'on screen' else ' · range: ' + src), chip='; '.join(chips), help=' · '.join(helps), tags='; '.join(tags),
            dname=('%s!%s' % (sheet, mp['name'])) if mp and sheet else '', dcell=di['cell'] if di else '', ddef=dd, dallowed=da, check=chk)
    # inflation / assumption notes
    for t in uniq(N.get('inflNote', []) + P.get('inflNote', [])): add(typ='Note', sec='Below the inputs', label=t)
    # result: gate, result, line, rows
    g = N.get('gate')
    if g: add(typ='Gate message', sec='Result', label=g['msg'], noplan=g['msg'] + (' · ' + g['also'] if g.get('also') else ''), plan=(P['gate']['msg'] if P.get('gate') else '(no gate: the result shows)'), behaviour='Shown until the choices the result needs are made; "Also still to add or choose" lists the rest.')
    pr, cr = P.get('result'), C.get('result')
    def match(label, val, nums):
        nl = norm(label)
        cands = [nm for nm, (dl, co) in rl.items() if norm(dl) == nl]
        if not cands: cands = [nm for nm, (dl, co) in rl.items() if norm(dl).startswith(nl) or nl.startswith(norm(dl))]
        if len(cands) == 1: return cands[0]
        v = num(val)
        if v is not None and nums:
            c2 = [nm for nm in rl if isinstance(nums.get(nm), (int, float)) and abs(nums[nm] - v) < 0.51]
            if len(c2) == 1: return c2[0]
        return None
    def devcols(name):
        if not name or not sheet: return {}
        return dict(dname='%s!%s' % (sheet, name), dcell=rl[name][1] if name in rl else '', dallowed='')
    if pr or cr:
        r = pr or cr
        nm = match(r['lbl'], r['val'], rec.get('numPlan') or rec.get('numChosen') or {})
        add(typ='Result', sec='Result', label=(pr or cr)['lbl'], noplan=(cr or {}).get('val', '(gated until a choice is made)'), plan=(pr or {}).get('val', '(gated)'), **devcols(nm))
        if (cr or {}).get('exres'): add(typ='Example-figures label (on result)', sec='Result', label=cr['exres'], noplan='Shown under the result', plan='Not shown', behaviour='Results made from example figures carry the label.')
        add(typ='Result line', sec='Result', label='Line under the result', noplan=(cr or {}).get('line') or '', plan=(pr or {}).get('line') or '')
        k = max(len((pr or {}).get('rows', [])), len((cr or {}).get('rows', [])))
        for i in range(k):
            a = (cr or {}).get('rows', [])[i] if i < len((cr or {}).get('rows', [])) else None
            b = (pr or {}).get('rows', [])[i] if i < len((pr or {}).get('rows', [])) else None
            lab = (b or a)[0]
            nm = match(lab, (b or a)[1], rec.get('numPlan') or {})
            add(typ='Sub-result row', sec='Result', label=lab, noplan=a[1] if a else '', plan=b[1] if b else '', **devcols(nm))
    # journey-spec 27.3: tools that follow the retirement age, read again with the retirement age set to 50
    E = rec.get('early')
    if E and E.get('result'):
        er = E['result']; base_rows = {(a, b) for a, b in ((pr or cr or {}).get('rows', []))}
        add(typ='Result (retiring at 50)', sec='Result', label=er['lbl'], plan=er['val'], source='The sample plan with the retirement age set to 50', behaviour='Same tool with the retirement age at 50: income stops at 50, the years before the pension (60) and the State Pension (66) are paid from savings, and the gap is shown.')
        m = re.search(r"Most pensions can't be taken before.*?savings and other income\.", er.get('line') or '')
        if m: add(typ='Note', sec='Result', label=m.group(0).strip(), plan='Shown when retiring at 50', noplan='Not shown', behaviour='Early-retirement warning: shown when the retirement age is before the pension access age (journey-spec 27.3). The access age is the Settings standard "Pension access age" (60, or 50 for some occupational schemes).')
        for lab, val in er.get('rows', []):
            if (lab, val) in base_rows: continue
            add(typ='Sub-result row (retiring at 50)', sec='Result', label=lab, plan=val, source='The sample plan with the retirement age set to 50')
    # buttons and their behaviour
    ap, an, ag = rec.get('addPlan'), rec.get('addNone'), rec.get('addGated')
    beh = []
    if ag and ag.get('toast'): beh.append('Choices missing: a message "%s"' % re.sub(r'^\S*\s*', '', ag['toast']) if False else 'Choices missing: a message "%s".' % ag['toast'].replace('✨', ''))
    if an and an.get('h2'): beh.append('No plan, choices made: a sheet "%s": "%s" Buttons: %s.' % (an['h2'], an['sub'], ', '.join('"%s"' % b for b in an['buttons'])))
    if ap: beh.append('With a plan: a sheet "%s": "%s" Buttons: %s. (Eyebrow "%s")' % (ap.get('h2'), ap.get('sub'), ', '.join('"%s"' % b for b in ap['buttons']), ap.get('eyebrow')))
    for b in N['buttons']:
        if b['action'] == 'calcadd': add(typ='Button', sec='Bottom', label=b['text'], behaviour=' '.join(beh), noplan='Tap', plan='Tap')
        else:
            e = rec.get('expert') or {}
            add(typ='Button', sec='Bottom', label=b['text'], noplan='Tap', plan='Tap', behaviour='Opens the Experts tab (screen %s "%s"), pre-set to the %s specialist: %s' % (e.get('id'), e.get('heading'), e.get('exp'), (e.get('what') or '')[:120]))
    if N.get('disc'): add(typ='Disclaimer', sec='Bottom', label=N['disc'])
    return rows


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('workbook'); ap.add_argument('--out'); ap.add_argument('--extract'); ap.add_argument('--shots')
    ap.add_argument('--no-recalc', action='store_true'); ap.add_argument('--recalc', default=DEFAULT_RECALC); ap.add_argument('--keep-work')
    a = ap.parse_args()
    work = a.keep_work or tempfile.mkdtemp(prefix='ui_sheets_')
    os.makedirs(work, exist_ok=True)
    ex_path, shots = a.extract, a.shots
    if not ex_path:
        out = os.path.join(work, 'extract'); env = dict(os.environ)
        env.setdefault('NODE_PATH', '/opt/node22/lib/node_modules')
        print('extracting from the live prototype ...', flush=True)
        subprocess.run(['node', os.path.join(HERE, 'extract_ui.js'), out, '--shots'], check=True, env=env)
        ex_path, shots = os.path.join(out, 'ui-extract.json'), os.path.join(out, 'shots')
    X = json.load(open(ex_path))
    shots = shots or os.path.join(os.path.dirname(ex_path), 'shots')
    calcmap = json.load(open(os.path.join(HERE, 'ui_calc_map.json')))
    out_path = a.out or a.workbook
    if a.out and os.path.abspath(a.out) != os.path.abspath(a.workbook): shutil.copyfile(a.workbook, a.out)
    wb = openpyxl.load_workbook(out_path); wbv = openpyxl.load_workbook(out_path, data_only=True)
    for t in [t for t in wb.sheetnames if t.startswith('UI ')]: del wb[t]
    dev = Dev(wb, wbv, calcmap)
    try:
        commit = subprocess.run(['git', '-C', REPO, 'rev-parse', '--short', 'HEAD'], capture_output=True, text=True).stdout.strip()
    except Exception:
        commit = ''
    g = wb.create_sheet('UI Guide'); c = wb.create_sheet('UI Calculators'); s = wb.create_sheet('UI Screens'); m = wb.create_sheet('UI Make my plan')
    for w in (g, c, s, m): w.sheet_properties.tabColor = TEAL

    # ------------------------------------------------------------ UI Calculators
    CH = ['Calculator', 'Developer sheet', 'Explore group', 'Calculator title', 'Question it answers', '#', 'Element type', 'Where on the screen', 'Label exactly as shown', 'Unit', 'Min', 'Max', 'Step',
          'Starts at WITHOUT a plan (as the screen shows it)', 'Starts at WITH a plan (Aoife sample)', 'Where the with-plan figure comes from', 'Choice / chip text', 'Help / hint / behaviour', 'Tags shown',
          'Developer name (sheet!name)', 'Developer cell', 'Developer default (example value)', 'Developer allowed range (validation)', 'Screen vs developer sheet']
    CW = [8, 22, 16, 24, 28, 5, 16, 15, 34, 7, 9, 9, 7, 24, 24, 30, 34, 46, 18, 30, 9, 11, 22, 40]
    c['A1'] = 'UI Calculators: one row per screen element of the 28 calculators, read from the live prototype (%s)' % X['meta']['prototype']; c['A1'].font = Font(bold=True, size=13, color=INK)
    c['A2'] = ('"Without a plan" = opened from Explore with no profile (the tool shows its own example figures, labelled). "With a plan" = the Aoife sample customer. Results "without a plan" are shown after the standards '
               'are chosen. Colours: see UI Guide. Developer columns come from the developer sheet of the same calculator in this workbook.')
    c['A2'].alignment = Alignment(wrap_text=True, vertical='top'); c.merge_cells('A2:X2'); c.row_dimensions[2].height = 32
    header(c, 4, CH, CW); c.freeze_panes = 'J5'
    r = 5; span = {}
    for rec in X['calcs']:
        first = r
        sheet = dev.sheet_for(rec['id'], rec['n'])
        for d in calc_rows(rec, dev, None):
            vals = [rec['n'], sheet or '', rec['group'], re.sub(r'^\S+\s', '', rec['none']['title']) if re.match(r'^\S+\s', rec['none']['title']) else rec['none']['title'], rec['q'], r - first + 1, d['typ'], d['sec'], d['label'], d['unit'], d['mn'], d['mx'], d['st'],
                    d['noplan'], d['plan'], d['source'], d['chip'], d['help'] if not d['behaviour'] else (d['behaviour'] + (' · ' + d['help'] if d['help'] else '')), d['tags'], d['dname'], d['dcell'], d['ddef'], d['dallowed'], d['check']]
            put(c, r, vals, kind=(7, kind_fill(d['typ'])))
            if sheet: link(c.cell(row=r, column=2), sheet, d['dcell'] or 'A1', sheet)
            r += 1
        span[rec['n']] = (first, r - 1)
    c.auto_filter.ref = 'A4:%s%d' % (get_column_letter(len(CH)), r - 1)
    calc_last = r - 1

    # ------------------------------------------------------------ UI Screens
    s['A1'] = 'UI Screens: each calculator as it looks in the app (390 px wide), without a plan and with the sample plan'; s['A1'].font = Font(bold=True, size=13, color=INK)
    s.column_dimensions['A'].width = 3; s.column_dimensions['B'].width = 58; s.column_dimensions['C'].width = 3; s.column_dimensions['D'].width = 58
    imgdir = os.path.join(work, 'png'); os.makedirs(imgdir, exist_ok=True)
    def quant(path):
        im = PILImage.open(path).convert('RGB'); q = im.quantize(256, method=PILImage.Quantize.MEDIANCUT, dither=PILImage.Dither.NONE)
        o = os.path.join(imgdir, os.path.basename(path)); q.save(o, optimize=True); return o, q.size
    r = 3; nimg = 0
    for rec in X['calcs']:
        n = rec['n']; f, l = span[n]
        t = s.cell(row=r, column=2, value='%s · %s · %s' % (n, rec['name'], rec['group'])); t.font = Font(bold=True, size=12, color='FFFFFF'); t.fill = fill(TEAL)
        s.cell(row=r, column=3).fill = fill(TEAL); s.cell(row=r, column=4).fill = fill(TEAL)
        link(s.cell(row=r, column=4), 'UI Calculators', 'A%d' % f, 'Rows %d to %d in UI Calculators' ' (%d elements)' % (f, l, l - f + 1)); s.cell(row=r, column=4).font = Font(bold=True, color='FFFFFF', underline='single')
        s.cell(row=r + 1, column=2, value='Without a plan (example figures)').font = Font(italic=True); s.cell(row=r + 1, column=4, value='With the sample plan (Aoife)').font = Font(italic=True)
        hmax = 0
        for col, key in ((2, 'shotNone'), (4, 'shotPlan')):
            p = os.path.join(shots, os.path.basename(rec[key])) if rec.get(key) else None
            if p and os.path.exists(p):
                o, (w, h) = quant(p); im = XlImage(o); s.add_image(im, '%s%d' % (get_column_letter(col), r + 2)); hmax = max(hmax, im.height); nimg += 1
        r += 2 + int(math.ceil(hmax / 20.0)) + 2
        if rec.get('shotEarly'):
            p = os.path.join(shots, os.path.basename(rec['shotEarly']))
            if os.path.exists(p):
                s.cell(row=r, column=2, value='%s retiring at 50 (journey-spec 27.3: the early-retirement warning and the years before the State Pension)' % n).font = Font(italic=True)
                o, (w, h) = quant(p); im = XlImage(o); s.add_image(im, 'B%d' % (r + 1)); nimg += 1
                r += 1 + int(math.ceil(im.height / 20.0)) + 2
    # ------------------------------------------------------------ UI Make my plan
    MH = ['Screen', 'Screen name', '#', 'Element type', 'Label / text exactly as shown', 'Unit', 'Min', 'Max', 'Step', 'Starts at / pre-filled with (Aoife sample)', 'Where the pre-fill comes from (tag on screen)', 'Choice / option texts',
          'Help / hint / behaviour', 'Example card (Not sure? See an example)', 'Picture (390 px)']
    MW = [10, 26, 4, 18, 46, 6, 8, 8, 6, 24, 28, 44, 44, 52, 56, 36]
    m['A1'] = 'UI Make my plan: one row per element of the plan-builder steps, the results and Your assumptions (from the live prototype)'; m['A1'].font = Font(bold=True, size=13, color=INK)
    m['A2'] = ('Steps 1 to 8: 1 About you & family, 2 Your goals, 3 Your timeline (and Confirm), 4 Your money profile, 5 Secure your account, 6 Your finances (six sections), 7 Check your details, 8 Results; then the journey-spec 27 states (gate card, Step 7 lists, free retirement age, Date has passed, early retirement, details missing). '
               'Pre-filled values are the sample customer\'s. A "Where the pre-fill comes from" tag is the tag the screen shows.')
    m['A2'].alignment = Alignment(wrap_text=True, vertical='top'); m.merge_cells('A2:O2'); m.row_dimensions[2].height = 32
    header(m, 4, MH, MW); m.freeze_panes = 'F5'
    r = 5; mp_span = {}; last_group = None
    for sc in X['plan']:
        if sc.get('group') and sc['group'] != last_group:
            if last_group is None:
                c0 = m.cell(row=r, column=2, value='Journey-spec 27 states: the screens as they look in each state (gate card, Step 7, free retirement age, "Date has passed", early retirement, details missing). Each block starts with a "Screen state" row.'); c0.font = Font(bold=True, size=12, color=INK); r += 1
            c1 = m.cell(row=r, column=2, value='State group: ' + sc['group']); c1.font = Font(bold=True, color='FFFFFF'); c1.fill = fill(TEAL); r += 1; last_group = sc['group']
        first = r
        if sc.get('note'):
            put(m, r, [sc['screen'], sc['name'], '', 'Screen state', 'State: ' + sc['name'], '', '', '', '', '', '', '', sc['note'], '', ''], kind=(4, kind_fill('Screen state'))); r += 1
        for i, e in enumerate(sc['rows'], 1):
            typ = e['type']
            ex = ' || '.join('%s: %s' % (k, X['examples'].get(k)) for k in (e.get('example') or []) if X['examples'].get(k))
            ch = '; '.join(uniq([o['t'] for o in (e.get('options') or [])] + [x['t'] for x in (e.get('chips') or [])]))
            helps = uniq(((e.get('help') or []) + (e.get('hint') or [])))
            beh = []
            if e.get('action') and typ in ('Button', 'Link', 'Input choice (chip)'): beh.append('action: %s%s' % (e['action'], ' ' + e['param'] if e.get('param') else ''))
            if e.get('disabled'): beh.append('off until the required choices are made')
            if e.get('parts') and len(e['parts']) > 1: beh.append(' | '.join(e['parts']))
            val = e.get('value')
            if typ.startswith('Input') and e.get('selected'): val = ', '.join(e['selected'])
            vals = [sc['screen'].split(' ')[0], sc['name'], i, typ, e.get('text') or e.get('label') or '', e.get('unit') or '', e.get('min') or '', e.get('max') or '', e.get('step') or '', val if typ.startswith('Input') else '',
                    '; '.join(e.get('tags') or []), ch, ' · '.join(helps + beh), ex, '']
            put(m, r, vals, kind=(4, kind_fill(typ))); r += 1
        last = r - 1
        if sc.get('shot'):
            p = os.path.join(shots, os.path.basename(sc['shot']))
            if os.path.exists(p):
                o, (w, h) = quant(p); im = XlImage(o); need = int(math.ceil(im.height / 20.0)) + 1
                m.add_image(im, 'O%d' % first); nimg += 1
                while r - first < need: put(m, r, ['', '', '', '', '', '', '', '', '', '', '', '', '', '', '']); r += 1
        mp_span[sc['screen']] = (first, r - 1)
    # Settings standards as the customer sees them
    r += 2
    m.cell(row=r, column=1, value='Your assumptions: the Settings standards as the customer sees them (wording and chip text)').font = Font(bold=True, size=12, color=INK); r += 1
    m.cell(row=r - 1, column=1).value = 'Your assumptions: the %d Settings standards as the customer sees them (wording and chip text). Tags in brackets, and the "Back to ..." button, appear once the customer has changed the figure (it then reads "Your choice").' % len(X['standards'])
    SH = ['Group', 'Standard (screen label)', '', 'Tag', 'Standard value', 'Unit', 'Chip text', 'Help text on the field', 'Wording of the standard', 'Source', 'As at', 'Verify before release', 'Developer name', 'Developer cell', 'Developer value', 'Name in the Settings sheet']
    for i, t in enumerate(SH, 1):
        cc = m.cell(row=r, column=i, value=t); cc.font = HEAD_FONT; cc.fill = fill(TEAL); cc.border = BORDER; cc.alignment = Alignment(wrap_text=True)
    r += 1
    setmap = {x['key']: x for x in X['settings']}
    st_ws = wb['Settings'] if 'Settings' in wb.sheetnames else None
    for sd in X['standards']:
        k = sd['key']; se = setmap.get(k)
        dn = 'Set_' + k; dcell = dval = ''
        if st_ws is not None and dn in wb.defined_names:
            try:
                sh, ref = list(wb.defined_names[dn].destinations)[0]; dcell = '%s!%s' % (sh, ref.replace('$', '')); dval = wbv[sh][ref.replace('$', '')].value
            except Exception:
                pass
        ac = sd.get('afterChange') or {}
        tags = uniq(sd['tags'] + [t + ' (after a change)' for t in ac.get('tags', [])])
        vals = [sd['group'], sd['label'], '', '; '.join(tags), sd['standard'], sd['unit'] or '', '; '.join(uniq(sd['chips'] + ac.get('back', []))), ' '.join(sd['help']), (se or {}).get('wording') or '', (se or {}).get('source') or '', (se or {}).get('asat') or '',
                'yes' if (se or {}).get('verify') else ('no' if se else ''), dn if dcell else '', dcell, dval if dval is not None else '', (se or {}).get('name') or '']
        put(m, r, vals); r += 1
    # ------------------------------------------------------------ UI Guide
    g.column_dimensions['A'].width = 30; g.column_dimensions['B'].width = 60; g.column_dimensions['C'].width = 40; g.column_dimensions['D'].width = 16
    g['A1'] = 'UI Guide: how to read the UI layer of this workbook'; g['A1'].font = Font(bold=True, size=14, color=INK)
    lines = [
        ('What this is', 'The developer sheets (C01 to C28, the Plan sheets, Tax engine, Settings and the others) hold the formulas. The four "UI" sheets (teal tabs, listed below with links) show the same 28 calculators, and the Make-my-plan journey with its journey-spec 27 states, as they look in the LifeMap application, so the UI/UX person and the developer work from one file.'),
        ('Where it comes from', 'Every label, range, start value, chip, help text and picture was read from the running prototype (%s%s) by tools/extract_ui.js with a real browser. Nothing was typed by hand. Re-run tools/build_ui_sheets.py on the final workbook to refresh it.' % (X['meta']['prototype'], (' at commit ' + commit) if commit else '')),
        ('How to read UI Calculators', 'One row per element of a screen, top to bottom. Filter column A for a calculator or column G for an element type. The "Developer ..." columns point to the cell in that calculator\'s developer sheet (click the sheet name to jump to the cell). The last column compares the screen with that cell.'),
        ('"Without a plan" and "With a plan"', 'Without a plan: the tool opened from Explore with no profile; it shows its own example figures, labelled "Example figures. Change them to yours." With a plan: the Aoife sample customer, whose own figures fill the inputs. Results without a plan are read after the standards are chosen.'),
        ('Developer vs screen', 'Slider ranges on screen are narrower than the developer\'s allowed ranges by design. The last column says "default same" or "DEFAULT DIFFERS" for the start value, and whether the slider range sits inside the developer\'s allowed range. A blank start on screen means the customer must choose (journey-spec §14); the developer sheet may carry an example default.'),
        ('Screenshots', 'UI Screens has two pictures per calculator (390 px wide, without a plan and with the sample plan) next to the row range in UI Calculators; a third picture (retiring at 50) for each tool that follows the retirement age. UI Make my plan has the picture of each step and of each journey-spec 27 state beside its rows.'),
        ('Journey-spec 27 states', 'UI Make my plan also shows the states added by journey-spec 27: the "To see your results we need N things" card on Results, Home and the Report (1 to 4 things, with and without a working partner) and its jump buttons; Step 7 "Choose 3 / 4 things" and "What we\'ve set for you (change any)" with the tags Set by LifeMap, Your choice, Assumed: add yours, Not chosen yet and the "Back to LifeMap\'s figure" button; the free retirement age (help text and the two calm notes, for the customer and the partner); "Date has passed" with "Change the date"; early-retirement results (income stops, bridge years, shortfall); and the "N details missing" banner with the Assumed rows. Each state starts with a "Screen state" row. The Settings standards table lists all %d standards, including "Pension access age" (From 60 (most pensions) / From 50 (some occupational schemes)).' % len(X['standards']))]
    r = 3
    for k, v in lines:
        g.cell(row=r, column=1, value=k).font = Font(bold=True); cc = g.cell(row=r, column=2, value=v); cc.alignment = Alignment(wrap_text=True, vertical='top'); g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        g.cell(row=r, column=1).alignment = Alignment(vertical='top'); g.row_dimensions[r].height = max(30, 15 * (len(v) // 100 + 1)); r += 1
    r += 1; g.cell(row=r, column=1, value='The four UI sheets (all in this workbook; click to open)').font = Font(bold=True, size=12, color=INK); r += 1
    for nm, ds in (('UI Guide', 'This sheet: how to read the layer, the colour legend, the screen order and links to every sheet.'), ('UI Calculators', 'One row per screen element of the 28 calculators, with the developer cell each maps to and a check against it.'),
                   ('UI Screens', 'A 390 px picture of each calculator, without a plan and with the sample plan (and retiring at 50 where it matters).'), ('UI Make my plan', 'One row per element of the eight steps, the results, the journey-spec 27 states and the Settings standards, with the pictures.')):
        link(g.cell(row=r, column=1), nm, 'A1', nm); d = g.cell(row=r, column=2, value=ds); d.alignment = Alignment(wrap_text=True, vertical='top'); g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); g.row_dimensions[r].height = 30; r += 1
    r += 1; g.cell(row=r, column=1, value='Colour legend (element type column)').font = Font(bold=True, size=12, color=INK); r += 1
    for name, kf, desc in LEGEND:
        cc = g.cell(row=r, column=1, value=name); cc.fill = fill(FILL[kf]); cc.border = BORDER; cc.alignment = Alignment(wrap_text=True, vertical='top')
        d = g.cell(row=r, column=2, value=desc); d.alignment = Alignment(wrap_text=True, vertical='top'); g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); g.row_dimensions[r].height = 30; r += 1
    r += 1; g.cell(row=r, column=1, value='Screen order in the app').font = Font(bold=True, size=12, color=INK); r += 1
    order = [('Explore', 'The Explore tab: "Focus on one area" tiles, the Calculators list in six groups, videos and tools.')]
    for gr in X['groups']:
        order.append(('Calculators · ' + gr['name'], ', '.join('%s %s' % (next(c['n'] for c in X['calcs'] if c['id'] == t), next(c['name'] for c in X['calcs'] if c['id'] == t)) for t in gr['tools'])))
    for i, (k, v) in enumerate(order):
        g.cell(row=r, column=1, value=k).font = Font(bold=True); cc = g.cell(row=r, column=2, value=v); cc.alignment = Alignment(wrap_text=True, vertical='top'); g.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        g.cell(row=r, column=1).alignment = Alignment(vertical='top', wrap_text=True); g.row_dimensions[r].height = max(18, 15 * (len(v) // 90 + 1)); r += 1
    r += 1; g.cell(row=r, column=1, value='Make my plan: steps 1 to 8, Results and the journey-spec 27 states').font = Font(bold=True, size=12, color=INK); r += 1
    header_row = ['Screen', 'Name', 'Rows in UI Make my plan', '']
    for i, t in enumerate(header_row, 1):
        cc = g.cell(row=r, column=i, value=t); cc.font = HEAD_FONT; cc.fill = fill(TEAL)
    r += 1
    for sc in X['plan']:
        a0, b0 = mp_span[sc['screen']]
        g.cell(row=r, column=1, value=sc['screen']); g.cell(row=r, column=2, value=sc['name']); link(g.cell(row=r, column=3), 'UI Make my plan', 'A%d' % a0, 'rows %d to %d' % (a0, b0)); r += 1
    r += 1; g.cell(row=r, column=1, value='Links to every sheet').font = Font(bold=True, size=12, color=INK); r += 1
    for i, t in enumerate(['Calculator', 'Developer sheet', 'UI rows', 'Picture'], 1):
        cc = g.cell(row=r, column=i, value=t); cc.font = HEAD_FONT; cc.fill = fill(TEAL)
    r += 1
    pos = {}
    rr = 3
    for rec in X['calcs']: pass
    for rec in X['calcs']:
        n = rec['n']; sheet = dev.sheet_for(rec['id'], n); a0, b0 = span[n]
        g.cell(row=r, column=1, value='%s %s' % (n, rec['name']))
        if sheet: link(g.cell(row=r, column=2), sheet, 'A1', sheet)
        link(g.cell(row=r, column=3), 'UI Calculators', 'A%d' % a0, 'rows %d to %d' % (a0, b0)); g.cell(row=r, column=4, value='see UI Screens'); r += 1
    for name in ('UI Guide', 'UI Calculators', 'UI Screens', 'UI Make my plan'):
        g.cell(row=r, column=1, value=name); link(g.cell(row=r, column=2), name, 'A1', name); r += 1
    for other in [t for t in wb.sheetnames if not t.startswith('UI ') and not re.match(r'C\d\d ', t)]:
        g.cell(row=r, column=1, value=other + ' (developer)'); link(g.cell(row=r, column=2), other, 'A1', other); r += 1
    wb.save(out_path)
    print('saved', out_path, '· UI Calculators rows 5 to %d · images %d' % (calc_last, nimg))
    if not a.no_recalc and os.path.exists(a.recalc):
        res = subprocess.run(['python3', a.recalc, out_path, '120'], capture_output=True, text=True); print(res.stdout.strip()[:400])
    elif not a.no_recalc:
        print('recalc.py not found: cached values were NOT restored (open in Excel or LibreOffice to recalculate)', file=sys.stderr)
    print('size %.1f MB' % (os.path.getsize(out_path) / 1e6))


if __name__ == '__main__':
    main()
